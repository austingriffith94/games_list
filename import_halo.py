"""Copy the 85 HaloRankMaster level rows from the workbook into halo_levels.

Usage: python import_halo.py <xlsm> <db>        (refuses to touch a non-empty halo_levels table)

Only the inputs are copied - game, level, order, the nine ratings and notes. Variable Score / Rank / Score
are derived on the fly (halo.py). A level whose nine ratings are all 0 is an unrated placeholder in the
workbook (Halo 5 and Halo Infinite) and is stored as NULL ratings, i.e. "unrated", not as a tier-F level.
"""
import sys, sqlite3, warnings
import openpyxl
from openpyxl.utils import range_boundaries
import halo

warnings.filterwarnings('ignore')


def read_levels(xlsm):
    wb = openpyxl.load_workbook(xlsm, data_only=True)
    ws = wb['HaloRank']
    c1, r1, c2, r2 = range_boundaries(ws.tables['HaloRankMaster'].ref)
    hdr = [str(ws.cell(r1, c).value) for c in range(c1, c2 + 1)]
    ix = {h: i for i, h in enumerate(hdr)}
    # the workbook's headers differ from the weights table's labels; ratings are taken by column *name* here,
    # and weights are paired to them by position later (see halo.py)
    want = ['Enemy', 'Weapon/Vehicle', 'Design/Set Piece', 'Replayability/Pacing', 'Difficulty', 'Visual', 'Music', 'Story', 'Enjoyment']
    out = []
    for r in range(r1 + 1, r2 + 1):
        row = [ws.cell(r, c).value for c in range(c1, c2 + 1)]
        if not row[ix['Game']]:
            continue
        ratings = [row[ix[h]] for h in want]
        ratings = [None if v is None else int(v) for v in ratings]
        if all(v in (None, 0) for v in ratings):
            ratings = [None] * 9                       # unrated placeholder
        elif any(v is None for v in ratings):
            raise ValueError(f'{row[ix["Game"]]} / {row[ix["Level"]]}: partially rated row {ratings}')
        note = row[ix['Notes']]
        out.append((int(row[ix['Series Order']]), row[ix['Game']], int(row[ix['Order']]), row[ix['Level']], *ratings,
                    str(note).strip() if note not in (None, '') else None))
    return out


def import_halo(xlsm, db):
    rows = read_levels(xlsm)
    con = sqlite3.connect(db)
    try:
        if con.execute('SELECT COUNT(*) FROM halo_levels').fetchone()[0]:
            raise SystemExit('halo_levels is not empty - nothing imported')
        cols = ['series_order', 'game', 'level_order', 'level', *halo.RATING_COLS, 'notes']
        con.executemany(f'INSERT INTO halo_levels({",".join(cols)}) VALUES ({",".join("?" * len(cols))})', rows)
        con.commit()
    finally:
        con.close()
    return len(rows), sum(1 for r in rows if r[4] is None)


if __name__ == '__main__':
    n, unrated = import_halo(sys.argv[1], sys.argv[2])
    print(f'imported {n} levels ({unrated} unrated)')
